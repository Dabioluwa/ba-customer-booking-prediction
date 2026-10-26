"""Build the one-slide manager summary from the analysis artefacts."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

matplotlib.use("Agg")

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs"
FIG_DIR = OUT_DIR / "figures"
DECK_PATH = ROOT / "Predicting_Customer_Booking_Summary.pptx"

# Colour system
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MID_BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT_BLUE = RGBColor(0x9D, 0xC3, 0xE6)
PALE_BLUE = RGBColor(0xDD, 0xEB, 0xF7)
TEAL = RGBColor(0x00, 0x8C, 0x95)
PALE_TEAL = RGBColor(0xE8, 0xF4, 0xF5)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x1E, 0x84, 0x49)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREY = RGBColor(0x6B, 0x72, 0x80)
DARK_GREY = RGBColor(0x33, 0x41, 0x55)
LIGHT_GREY = RGBColor(0xE2, 0xE8, 0xF0)
OFF_WHITE = RGBColor(0xF7, 0xF9, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Segoe UI"

# ---------------------------------------------------------------------------
# Load results
# ---------------------------------------------------------------------------
with open(OUT_DIR / "metrics.json", encoding="utf-8") as handle:
    metrics = json.load(handle)

grouped = pd.read_csv(OUT_DIR / "grouped_feature_importance.csv")
ranking = pd.read_csv(OUT_DIR / "targeting_lift.csv")

holdout = metrics["holdout"]
cv = metrics["cv"]
dataset = metrics["dataset"]
top10 = ranking.loc[ranking["target_fraction"].round(2) == 0.10].iloc[0]
top20 = ranking.loc[ranking["target_fraction"].round(2) == 0.20].iloc[0]

# ---------------------------------------------------------------------------
# Slide-specific driver chart
# ---------------------------------------------------------------------------
plot_data = grouped.head(7).iloc[::-1].copy()
plot_values = plot_data["importance_mean"] * 100

fig, ax = plt.subplots(figsize=(9.6, 3.88), facecolor="none")
fig.patch.set_alpha(0)
ax.set_facecolor("none")
colors = ["#2E75B6" if value < 5 else "#1F4E79" for value in plot_values]
colors[-1] = "#008C95"
colors[-2] = "#008C95"

bars = ax.barh(
    plot_data["driver"],
    plot_values,
    color=colors,
    height=0.62,
    zorder=3,
)
for bar, value in zip(bars, plot_values):
    ax.text(
        bar.get_width() + 0.12,
        bar.get_y() + bar.get_height() / 2,
        f"{value:.1f}",
        va="center",
        ha="left",
        fontsize=12,
        fontweight="bold",
        color="#0B1F3A",
    )

ax.set_xlabel("Drop in holdout ROC-AUC (percentage points)", fontsize=11, color="#334155", labelpad=8)
ax.set_xlim(0, 12.6)
ax.set_xticks([0, 2, 4, 6, 8, 10, 12])
ax.tick_params(axis="x", labelsize=10, colors="#475569", length=0)
ax.tick_params(axis="y", labelsize=12, colors="#0B1F3A", length=0, pad=6)
ax.grid(axis="x", color="#E2E8F0", linewidth=0.9, zorder=0)
ax.set_axisbelow(True)
for spine in ("top", "right", "left"):
    ax.spines[spine].set_visible(False)
ax.spines["bottom"].set_color("#CBD5E1")
ax.spines["bottom"].set_linewidth(0.8)

fig.subplots_adjust(left=0.31, right=0.97, top=0.97, bottom=0.19)
chart_path = FIG_DIR / "slide_driver_importance.png"
fig.savefig(chart_path, dpi=260, transparent=True)
plt.close(fig)

# ---------------------------------------------------------------------------
# Presentation helpers
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Background
background = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
)
background.fill.solid()
background.fill.fore_color.rgb = OFF_WHITE
background.line.fill.background()
background.shadow.inherit = False


def add_rect(
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    line: RGBColor | None = None,
    shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
    radius: float | None = None,
):
    rect = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = fill
    if line is None:
        rect.line.fill.background()
    else:
        rect.line.color.rgb = line
        rect.line.width = Pt(0.75)
    rect.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        rect.adjustments[0] = radius
    return rect


def add_text(
    x: float,
    y: float,
    w: float,
    h: float,
    text: str = "",
    size: float = 10,
    color: RGBColor = DARK_GREY,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    italic: bool = False,
    spacing: float = 1.0,
    margin: float = 0.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0)
    frame.margin_bottom = Inches(0)
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return box


def add_rich_text(
    x: float,
    y: float,
    w: float,
    h: float,
    runs: list[tuple[str, float, RGBColor, bool, bool]],
    align: PP_ALIGN = PP_ALIGN.LEFT,
    spacing: float = 1.0,
    space_after: float = 0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    paragraph.space_after = Pt(space_after)
    for text, size, color, bold, italic in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
    return box


# ---------------------------------------------------------------------------
# Header
# ---------------------------------------------------------------------------
add_rect(0, 0, 13.333, 1.03, NAVY)
add_rect(0, 1.03, 13.333, 0.045, TEAL)

add_text(
    0.45,
    0.17,
    10.9,
    0.52,
    "Customer booking prediction: geography is the strongest signal",
    24,
    WHITE,
    bold=True,
)
add_text(
    0.47,
    0.70,
    10.5,
    0.25,
    "Random Forest model on 49,281 British Airways booking enquiries  |  5-fold cross-validation + holdout evaluation",
    10.5,
    LIGHT_BLUE,
)

add_text(11.0, 0.33, 1.9, 0.22, "CUSTOMER ANALYTICS", 8.5, LIGHT_BLUE, bold=True, align=PP_ALIGN.RIGHT)
add_text(11.0, 0.55, 1.9, 0.22, "October 2026", 8.5, RGBColor(0xBF, 0xD4, 0xEA), align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------------------
# KPI cards
# ---------------------------------------------------------------------------
card_y = 1.22
card_h = 1.05
card_w = 3.0
card_xs = [0.45, 3.593, 6.735, 9.878]
kpis = [
    (
        f"{holdout['roc_auc']:.2f}",
        "ROC-AUC",
        f"Holdout  |  CV {cv['roc_auc']['mean']:.3f} ± {cv['roc_auc']['std']:.3f}",
        BLUE,
    ),
    (
        f"{holdout['pr_auc']:.2f}",
        "PR-AUC",
        f"{holdout['pr_auc'] / holdout['prevalence']:.1f}x the {holdout['prevalence']:.0%} base rate",
        TEAL,
    ),
    (
        f"{top10['lift_vs_random']:.1f}x",
        "LIFT AT TOP 10%",
        "vs contacting a random 10%",
        ORANGE,
    ),
    (
        f"{top10['share_of_all_bookings']:.0%}",
        "BOOKINGS CAPTURED",
        "by targeting the top 10%",
        GREEN,
    ),
]

for x, (value, label, sub, accent) in zip(card_xs, kpis):
    add_rect(x, card_y, card_w, card_h, WHITE, line=LIGHT_GREY)
    add_rect(x, card_y, card_w, 0.065, accent)
    add_text(x + 0.18, card_y + 0.12, 1.15, 0.46, value, 26, NAVY, bold=True)
    add_text(x + 0.18, card_y + 0.59, card_w - 0.3, 0.20, label, 8, GREY, bold=True)
    add_text(x + 0.18, card_y + 0.79, card_w - 0.3, 0.20, sub, 8, GREY)

# ---------------------------------------------------------------------------
# Chart card
# ---------------------------------------------------------------------------
chart_card_x = 0.45
chart_card_y = 2.44
chart_card_w = 7.55
chart_card_h = 3.70
add_rect(chart_card_x, chart_card_y, chart_card_w, chart_card_h, WHITE, line=LIGHT_GREY)
add_text(
    chart_card_x + 0.20,
    chart_card_y + 0.15,
    6.4,
    0.26,
    "What drives a booking?",
    14,
    NAVY,
    bold=True,
)
add_text(
    chart_card_x + 0.20,
    chart_card_y + 0.43,
    7.1,
    0.32,
    "Drop in holdout ROC-AUC when each driver is shuffled. Other features add <0.1 pts each.",
    8.2,
    GREY,
)

chart_pic = slide.shapes.add_picture(
    str(chart_path),
    Inches(chart_card_x + 0.06),
    Inches(chart_card_y + 0.62),
    width=Inches(chart_card_w - 0.12),
)
chart_pic._element.nvPicPr.cNvPr.set(
    "descr",
    "Horizontal bar chart of permutation importance. Booking origin 11.0 percentage points, "
    "flight route 7.3, optional extras 1.8, length of stay 1.1, sales channel 0.7, "
    "flight duration 0.5, purchase lead time 0.5.",
)

# ---------------------------------------------------------------------------
# Insight panel
# ---------------------------------------------------------------------------
panel_x = 8.15
panel_y = 2.44
panel_w = 4.73
panel_h = 3.70
add_rect(panel_x, panel_y, panel_w, panel_h, RGBColor(0xF0, 0xF5, 0xFA), line=RGBColor(0xDC, 0xE5, 0xF0))
add_rect(panel_x, panel_y, 0.07, panel_h, TEAL)

add_text(
    panel_x + 0.24,
    panel_y + 0.15,
    panel_w - 0.4,
    0.24,
    "WHAT THE MODEL TELLS US",
    9,
    TEAL,
    bold=True,
)

insights = [
    (
        "1",
        "Market and route drive ~80% of measured importance.",
        "Conversion is ~5% in Australia/New Zealand vs ~35% in Malaysia. The score mostly separates markets before it separates individuals.",
        MID_BLUE,
    ),
    (
        "2",
        "Intent signals add usable, but smaller, lift.",
        "All three extras: 18.6% vs 10.7% with none. Short breaks (4-7 days): 19.5% vs 10.1% for 15-30 days. Internet 15.5% vs Mobile 11.0%.",
        TEAL,
    ),
    (
        "3",
        "A useful ranking, not a guarantee.",
        f"PR-AUC {holdout['pr_auc']:.2f} vs {holdout['prevalence']:.2f} base rate; at the 0.5 threshold the model finds {holdout['recall']:.0%} of bookers with {holdout['precision']:.0%} precision.",
        ORANGE,
    ),
]

block_y = panel_y + 0.46
for number, headline, body, accent in insights:
    add_rect(panel_x + 0.24, block_y + 0.02, 0.28, 0.28, accent, shape=MSO_SHAPE.OVAL)
    add_text(
        panel_x + 0.24,
        block_y + 0.025,
        0.28,
        0.25,
        number,
        9,
        WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        panel_x + 0.61,
        block_y,
        panel_w - 0.78,
        0.26,
        headline,
        9.5,
        NAVY,
        bold=True,
    )
    add_text(
        panel_x + 0.61,
        block_y + 0.27,
        panel_w - 0.78,
        0.55,
        body,
        8.9,
        DARK_GREY,
        spacing=1.0,
    )
    block_y += 0.82

# Bottom-line note inside the panel
add_rect(panel_x + 0.24, panel_y + 3.00, panel_w - 0.46, 0.54, RGBColor(0xE2, 0xEC, 0xF7))
add_text(
    panel_x + 0.36,
    panel_y + 3.07,
    panel_w - 0.66,
    0.42,
    f"Accuracy is {holdout['accuracy']:.0%} - below the 85% 'never books' baseline. Precision at the top of the ranking is the metric that matters.",
    8.0,
    NAVY,
    italic=True,
    spacing=1.0,
)

# ---------------------------------------------------------------------------
# Recommendation bar
# ---------------------------------------------------------------------------
bar_y = 6.26
add_rect(0.45, bar_y, 12.43, 0.74, PALE_TEAL, line=RGBColor(0xC8, 0xE4, 0xE6))
add_rect(0.45, bar_y, 0.08, 0.74, TEAL)
add_text(0.68, bar_y + 0.10, 1.6, 0.20, "RECOMMENDATION", 8.5, TEAL, bold=True)
add_text(
    0.68,
    bar_y + 0.28,
    12.0,
    0.19,
    "Rank enquiries and target the top 10-20% for 2.5-2.9x lift; A/B test incremental bookings before scaling.",
    9.2,
    NAVY,
)
add_text(
    0.68,
    bar_y + 0.47,
    12.0,
    0.19,
    "Localise offers by origin/route and short breaks; retrain quarterly with price and seasonality data.",
    9.2,
    NAVY,
)

# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------
add_text(
    0.45,
    7.10,
    9.7,
    0.30,
    f"Holdout (n={dataset['test_rows']:,}): ROC-AUC {holdout['roc_auc']:.3f}  |  PR-AUC {holdout['pr_auc']:.3f}  |  recall {holdout['recall']:.2f}  |  precision {holdout['precision']:.2f} at 0.50.   "
    f"CV: ROC-AUC {cv['roc_auc']['mean']:.3f} ± {cv['roc_auc']['std']:.3f}.",
    7.0,
    GREY,
)
add_text(
    0.45,
    7.27,
    10.9,
    0.20,
    "Data: 50,000 rows, 719 exact duplicates removed. 80/20 stratified split; 5-fold CV; 600-tree Random Forest; out-of-fold target encoding for route and origin. Importance = grouped drop in ROC-AUC when a driver is shuffled.",
    6.5,
    RGBColor(0x94, 0xA3, 0xB8),
)
add_text(12.25, 7.22, 0.65, 0.20, "1", 7.5, GREY, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Speaker notes and document properties
# ---------------------------------------------------------------------------
notes_frame = slide.notes_slide.notes_text_frame
notes_frame.text = (
    "PURPOSE: one-slide summary of the customer-booking propensity model.\n\n"
    "DATA: 50,000 British Airways booking enquiries, 14 columns, no missing values. "
    "719 exact duplicate rows were removed, leaving 49,281 rows; 15.0% completed a booking. "
    "An 80/20 stratified split was used (39,424 train / 9,857 holdout).\n\n"
    "FEATURES: num_passengers, purchase_lead, length_of_stay, flight_hour (plus sine/cosine), "
    "flight_duration, flight_day_num, is_weekend, extras_count, the three optional-extra flags, "
    "sales_channel, trip_type, route and booking_origin. Route and booking_origin were "
    "target-encoded out-of-fold inside the model pipeline.\n\n"
    "MODEL: RandomForestClassifier, 600 trees, min_samples_leaf=10, max_features='sqrt', "
    "class_weight='balanced_subsample'. 5-fold stratified CV on the training set. "
    "CV ROC-AUC 0.795 +/- 0.003, PR-AUC 0.391 +/- 0.011, F1 0.451. "
    "Holdout ROC-AUC 0.788, PR-AUC 0.367 (vs 0.150 prevalence), accuracy 76.8%, recall 60.6%, "
    "precision 34.4% at the 0.5 threshold.\n\n"
    "BUSINESS RESULT: the model is a ranking tool. Contacting the top 10% captures 28.6% of all "
    "bookings (2.86x lift); the top 20% captures 50.3% (2.51x lift). Top decile precision is 42.8%, "
    "so most contacted customers still will not book - it is a prioritisation aid, not a guarantee.\n\n"
    "DRIVERS: grouped permutation importance on the holdout set. Booking origin is 11.0 percentage "
    "points of ROC-AUC, route 7.3, optional extras 1.8, length of stay 1.1, sales channel 0.7, "
    "flight duration 0.5, purchase lead time 0.5. Origin and route together are about 80% of "
    "measured importance, so the score mostly separates markets before it separates individuals.\n\n"
    "CAVEATS: (1) The optional-extra fields may be captured late in the funnel; without them the "
    "holdout ROC-AUC is 0.781 and PR-AUC 0.352. Confirm timing before production use. (2) The "
    "analysis is predictive, not causal. (3) Price, seasonality/travel dates, repeat-customer "
    "history and abandoned-search data were not available. (4) Retrain and monitor drift before "
    "reuse.\n\n"
    "RECOMMENDATION: deploy as a ranked marketing score, target the top 10-20%, A/B test incremental "
    "bookings against random contact, localise creative by origin/route and short breaks, and "
    "retrain quarterly with price and seasonality data."
)

prs.core_properties.title = "Predicting customer buying behaviour - model summary"
prs.core_properties.subject = "British Airways booking propensity model"
prs.core_properties.author = "Customer Analytics"
prs.core_properties.keywords = "customer booking, random forest, propensity model, ROC-AUC, permutation importance"

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
prs.save(DECK_PATH)
print(f"Saved {DECK_PATH}")

# Basic layout audit
for shape in slide.shapes:
    if shape.left < 0 or shape.top < 0:
        print(f"WARNING: shape outside slide: {shape.name}")
    if shape.left + shape.width > prs.slide_width + Emu(1):
        print(f"WARNING: shape exceeds right edge: {shape.name}")
    if shape.top + shape.height > prs.slide_height + Emu(1):
        print(f"WARNING: shape exceeds bottom edge: {shape.name}")
